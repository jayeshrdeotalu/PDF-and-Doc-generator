from docx import Document

def fill_placeholders(template_path, output_path, replacements):
    # Load the Word document
    doc = Document(template_path)
    
    # Loop through each paragraph in the document
    for paragraph in doc.paragraphs:
        for key, value in replacements.items():
            # Replace placeholder with actual value
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)
    
    # Save the updated document
    doc.save(output_path)

# Example usage
template_path = r'/home/om/Desktop/auto maker/autofill_temp.docx'  # Path to your Word template file
output_path = r'C:\Users\Jayesh B\Desktop\Auto maker\test_1.docx'      # Path to save the updated file
replacements = {
    '{NAME}': 'John Doe',
    '{STID}': '2024-08-12',
    '{EMAIL}': 'Software Intern',
    '{PHNO}': 'Software Intern',
    '{FOS}': 'Software Intern',
    '{YOS}': 'Software Intern',
    '{INT_JOI}': 'Software Intern',
    '{ROBO_F}': 'Software Intern',
    '{LEADER}': 'Software Intern',
    '{EXP}': 'Software Intern',
    '{QUE}': 'Software Intern',
    
}

fill_placeholders(template_path, output_path, replacements)
